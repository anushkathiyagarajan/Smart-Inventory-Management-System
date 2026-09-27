from flask import Flask, jsonify, send_file, send_from_directory
from pathlib import Path
from io import BytesIO
import pandas as pd
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

app = Flask(__name__)

BASE_DIR = Path(__file__).resolve().parent.parent
DATA_DIR = BASE_DIR / "data"
FRONTEND_DIR = BASE_DIR / "frontend"

DATA_FILE = DATA_DIR / "inventory_data.csv"
FORECAST_FILE = DATA_DIR / "demand_forecast.csv"
MODEL_FILE = DATA_DIR / "forecast_model_comparison.csv"


# ============================================================
# LOAD DATA
# ============================================================

def load_inventory_data():

    if not DATA_FILE.exists():
        raise FileNotFoundError(
            f"Inventory dataset not found: {DATA_FILE}"
        )

    df = pd.read_csv(DATA_FILE)

    df["Date"] = pd.to_datetime(
        df["Date"],
        errors="coerce"
    )

    numeric_columns = [
        "Inventory Level",
        "Units Sold",
        "Units Ordered",
        "Price",
        "Discount",
        "Competitor Pricing",
        "Demand"
    ]

    for column in numeric_columns:

        if column in df.columns:

            df[column] = pd.to_numeric(
                df[column],
                errors="coerce"
            ).fillna(0)

    # --------------------------------------------------------
    # STOCK COVERAGE
    # --------------------------------------------------------

    df["Stock Coverage"] = np.where(
        df["Demand"] > 0,
        df["Inventory Level"] / df["Demand"],
        0
    )

    # --------------------------------------------------------
    # STOCK STATUS
    # --------------------------------------------------------

    df["Stock Status"] = "Healthy Stock"

    df.loc[
        df["Stock Coverage"] < 1,
        "Stock Status"
    ] = "Low Stock"

    df.loc[
        df["Inventory Level"] == 0,
        "Stock Status"
    ] = "Out of Stock"

    # --------------------------------------------------------
    # INVENTORY VALUE
    # --------------------------------------------------------

    df["Inventory Value"] = (
        df["Inventory Level"] *
        df["Price"]
    )

    # --------------------------------------------------------
    # REORDER
    # --------------------------------------------------------

    df["Reorder Level"] = (
        df["Demand"] * 1.2
    ).round().astype(int)

    df["Reorder Quantity"] = (
        df["Reorder Level"] -
        df["Inventory Level"]
    ).clip(lower=0).round().astype(int)

    # --------------------------------------------------------
    # INVENTORY CONDITION
    # --------------------------------------------------------

    df["Inventory Condition"] = "Normal"

    df.loc[
        df["Stock Coverage"] >= 6,
        "Inventory Condition"
    ] = "Overstocked"

    # --------------------------------------------------------
    # RECOMMENDATION
    # --------------------------------------------------------

    df["Smart Recommendation"] = "No Action"

    df.loc[
        df["Stock Status"] == "Out of Stock",
        "Smart Recommendation"
    ] = "Urgent Restock"

    df.loc[
        (
            (df["Stock Status"] == "Low Stock") &
            (df["Reorder Quantity"] > 0)
        ),
        "Smart Recommendation"
    ] = "Restock Soon"

    df.loc[
        (
            (df["Inventory Condition"] == "Overstocked") &
            (df["Stock Status"] == "Healthy Stock")
        ),
        "Smart Recommendation"
    ] = "Overstock - Monitor"

    # --------------------------------------------------------
    # TIME FEATURES
    # --------------------------------------------------------

    df["Year"] = df["Date"].dt.year
    df["Month"] = df["Date"].dt.month
    df["Month Name"] = df["Date"].dt.strftime("%b")
    df["Year-Month"] = (
        df["Date"]
        .dt.to_period("M")
        .astype(str)
    )

    return df


def load_forecast():

    if not FORECAST_FILE.exists():
        return pd.DataFrame()

    return pd.read_csv(
        FORECAST_FILE
    )


def load_model_comparison():

    if not MODEL_FILE.exists():
        return pd.DataFrame()

    return pd.read_csv(
        MODEL_FILE
    )


# ============================================================
# SUMMARY
# ============================================================

def get_summary(df):

    product_sales = (
        df.groupby("Product ID")["Units Sold"]
        .sum()
        .sort_values(
            ascending=False
        )
    )

    category_sales = (
        df.groupby("Category")["Units Sold"]
        .sum()
        .sort_values(
            ascending=False
        )
    )

    product_analysis = (
        df.groupby("Product ID")
        .agg(
            Total_Sales=(
                "Units Sold",
                "sum"
            ),
            Total_Inventory=(
                "Inventory Level",
                "sum"
            )
        )
    )

    product_analysis[
        "Sell_Through_Rate"
    ] = (
        product_analysis["Total_Sales"] /
        (
            product_analysis["Total_Sales"] +
            product_analysis["Total_Inventory"]
        ).replace(0, 1)
    ) * 100

    return {

        "total_products":
            int(df["Product ID"].nunique()),

        "total_categories":
            int(df["Category"].nunique()),

        "available_stock":
            int(df["Inventory Level"].sum()),

        "inventory_value":
            float(df["Inventory Value"].sum()),

        "low_stock":
            int(
                (
                    df["Stock Status"] ==
                    "Low Stock"
                ).sum()
            ),

        "out_of_stock":
            int(
                (
                    df["Stock Status"] ==
                    "Out of Stock"
                ).sum()
            ),

        "overstocked":
            int(
                (
                    df["Inventory Condition"] ==
                    "Overstocked"
                ).sum()
            ),

        "healthy_stock":
            int(
                (
                    df["Stock Status"] ==
                    "Healthy Stock"
                ).sum()
            ),

        "top_product":
            str(
                product_sales.index[0]
            ),

        "top_product_sales":
            int(
                product_sales.iloc[0]
            ),

        "top_category":
            str(
                category_sales.index[0]
            ),

        "top_category_sales":
            int(
                category_sales.iloc[0]
            ),

        "fast_moving":
            int(
                (
                    product_analysis[
                        "Sell_Through_Rate"
                    ] >= 30
                ).sum()
            ),

        "moderate_moving":
            int(
                (
                    (
                        product_analysis[
                            "Sell_Through_Rate"
                        ] >= 20
                    ) &
                    (
                        product_analysis[
                            "Sell_Through_Rate"
                        ] < 30
                    )
                ).sum()
            ),

        "slow_moving":
            int(
                (
                    product_analysis[
                        "Sell_Through_Rate"
                    ] < 20
                ).sum()
            )
    }


# ============================================================
# FORECAST
# ============================================================

def get_forecast():

    df = load_forecast()

    if df.empty:

        return {
            "period": "N/A",
            "demand": 0,
            "method": "N/A",
            "data": []
        }

    rows = []

    for _, row in df.iterrows():

        rows.append({

            "period":
                str(
                    row.get(
                        "Forecast Period",
                        ""
                    )
                ),

            "demand":
                float(
                    row.get(
                        "Forecast Demand",
                        0
                    )
                ),

            "method":
                str(
                    row.get(
                        "Forecast Method",
                        ""
                    )
                )
        })

    return {

        "period":
            rows[0]["period"],

        "demand":
            rows[0]["demand"],

        "method":
            rows[0]["method"],

        "data":
            rows
    }

@app.route("/")
def home():

    html_file = FRONTEND_DIR / "index.html"

    if not html_file.exists():
        return "index.html not found", 404

    html = html_file.read_text(
        encoding="utf-8"
    )

    df = load_inventory_data()

    summary = get_summary(df)

    forecast = get_forecast()

    replacements = {

        "{{TOTAL_PRODUCTS}}":
            f"{summary['total_products']:,}",

        "{{TOTAL_CATEGORIES}}":
            f"{summary['total_categories']:,}",

        "{{AVAILABLE_STOCK}}":
            f"{summary['available_stock']:,}",

        "{{INVENTORY_VALUE}}":
            f"₹{summary['inventory_value']:,.0f}",

        "{{LOW_STOCK}}":
            f"{summary['low_stock']:,}",

        "{{OUT_OF_STOCK}}":
            f"{summary['out_of_stock']:,}",

        "{{OVERSTOCKED}}":
            f"{summary['overstocked']:,}",

        "{{HEALTHY_STOCK}}":
            f"{summary['healthy_stock']:,}",

        "{{TOP_PRODUCT}}":
            summary["top_product"],

        "{{TOP_PRODUCT_SALES}}":
            f"{summary['top_product_sales']:,}",

        "{{TOP_CATEGORY}}":
            summary["top_category"],

        "{{TOP_CATEGORY_SALES}}":
            f"{summary['top_category_sales']:,}",

        "{{FAST_MOVING}}":
            f"{summary['fast_moving']:,}",

        "{{MODERATE_MOVING}}":
            f"{summary['moderate_moving']:,}",

        "{{SLOW_MOVING}}":
            f"{summary['slow_moving']:,}",

        "{{FORECAST_PERIOD}}":
            forecast["period"],

        "{{FORECAST_DEMAND}}":
            f"{forecast['demand']:,.0f}",

        "{{FORECAST_METHOD}}":
            forecast["method"]
    }

    for key, value in replacements.items():

        html = html.replace(
            key,
            str(value)
        )

    return html

# ============================================================
# DASHBOARD API
# ============================================================

@app.route("/api/dashboard")
def dashboard():

    try:

        df = load_inventory_data()

        # ----------------------------------------------------
        # CATEGORY
        # ----------------------------------------------------

        category = (
            df.groupby("Category")
            .agg(
                Inventory_Level=(
                    "Inventory Level",
                    "sum"
                ),
                Units_Sold=(
                    "Units Sold",
                    "sum"
                ),
                Demand=(
                    "Demand",
                    "sum"
                ),
                Inventory_Value=(
                    "Inventory Value",
                    "sum"
                )
            )
            .reset_index()
        )

        category_data = []

        for _, row in category.iterrows():

            category_data.append({

                "category":
                    str(row["Category"]),

                "inventory":
                    int(row["Inventory_Level"]),

                "sales":
                    int(row["Units_Sold"]),

                "demand":
                    int(row["Demand"]),

                "value":
                    float(row["Inventory_Value"])
            })

        # ----------------------------------------------------
        # PRODUCTS
        # ----------------------------------------------------

        products = (
            df.groupby("Product ID")
            .agg(
                Total_Inventory=(
                    "Inventory Level",
                    "sum"
                ),
                Total_Sales=(
                    "Units Sold",
                    "sum"
                ),
                Total_Demand=(
                    "Demand",
                    "sum"
                )
            )
            .reset_index()
            .sort_values(
                "Total_Sales",
                ascending=False
            )
        )

        product_data = []

        for _, row in products.iterrows():

            product_data.append({

                "product":
                    str(row["Product ID"]),

                "inventory":
                    int(row["Total_Inventory"]),

                "sales":
                    int(row["Total_Sales"]),

                "demand":
                    int(row["Total_Demand"])
            })

        # ----------------------------------------------------
        # MONTHLY
        # ----------------------------------------------------

        monthly = (
            df.groupby("Year-Month")
            .agg(
                Inventory_Level=(
                    "Inventory Level",
                    "mean"
                ),
                Units_Sold=(
                    "Units Sold",
                    "sum"
                ),
                Demand=(
                    "Demand",
                    "sum"
                )
            )
            .reset_index()
            .sort_values(
                "Year-Month"
            )
        )

        monthly_data = []

        for _, row in monthly.iterrows():

            monthly_data.append({

                "month":
                    str(row["Year-Month"]),

                "inventory":
                    float(
                        row["Inventory_Level"]
                    ),

                "sales":
                    int(
                        row["Units_Sold"]
                    ),

                "demand":
                    int(
                        row["Demand"]
                    )
            })

        return jsonify({

            "success": True,

            "summary":
                get_summary(df),

            "category":
                category_data,

            "products":
                product_data,

            "monthly":
                monthly_data,

            "forecast":
                get_forecast()
        })

    except Exception as error:

        return jsonify({

            "success": False,

            "error": str(error)

        }), 500


# ============================================================
# DATA API
# ============================================================

@app.route("/api/data")
def api_data():

    try:

        df = load_inventory_data()

        display_df = df.head(
            1000
        ).copy()

        display_df["Date"] = (
            display_df["Date"]
            .dt.strftime(
                "%Y-%m-%d"
            )
        )

        display_df = display_df.replace(
            {np.nan: None}
        )

        return jsonify({

            "success": True,

            "rows":
                display_df.to_dict(
                    orient="records"
                ),

            "total_rows":
                int(len(df))
        })

    except Exception as error:

        return jsonify({

            "success": False,

            "error": str(error)

        }), 500


# ============================================================
# EXPORT CSV
# ============================================================

@app.route("/api/export-data")
def export_data():

    try:

        df = load_inventory_data()

        output = BytesIO()

        df.to_csv(
            output,
            index=False
        )

        output.seek(0)

        return send_file(

            output,

            mimetype="text/csv",

            as_attachment=True,

            download_name=
            "stockwise_inventory_data.csv"
        )

    except Exception as error:

        return jsonify({

            "success": False,

            "error": str(error)

        }), 500


# ============================================================
# EXCEL REPORT
# ============================================================

def create_excel_report(df):

    workbook = Workbook()

    # --------------------------------------------------------
    # SUMMARY
    # --------------------------------------------------------

    sheet = workbook.active

    sheet.title = "Inventory Summary"

    summary = get_summary(df)

    sheet.append([
        "Metric",
        "Value"
    ])

    rows = [

        [
            "Total Products",
            summary["total_products"]
        ],

        [
            "Total Categories",
            summary["total_categories"]
        ],

        [
            "Available Stock",
            summary["available_stock"]
        ],

        [
            "Inventory Value",
            summary["inventory_value"]
        ],

        [
            "Low Stock Records",
            summary["low_stock"]
        ],

        [
            "Out of Stock Records",
            summary["out_of_stock"]
        ],

        [
            "Overstocked Records",
            summary["overstocked"]
        ],

        [
            "Healthy Stock Records",
            summary["healthy_stock"]
        ],

        [
            "Top Selling Product",
            summary["top_product"]
        ],

        [
            "Top Performing Category",
            summary["top_category"]
        ]
    ]

    for row in rows:
        sheet.append(row)

    # --------------------------------------------------------
    # CATEGORY
    # --------------------------------------------------------

    category_sheet = workbook.create_sheet(
        "Category Analysis"
    )

    category = (
        df.groupby("Category")
        .agg(
            Inventory_Level=(
                "Inventory Level",
                "sum"
            ),
            Units_Sold=(
                "Units Sold",
                "sum"
            ),
            Demand=(
                "Demand",
                "sum"
            ),
            Inventory_Value=(
                "Inventory Value",
                "sum"
            )
        )
        .reset_index()
    )

    category_sheet.append(
        list(category.columns)
    )

    for row in category.itertuples(
        index=False,
        name=None
    ):

        category_sheet.append(
            list(row)
        )

    # --------------------------------------------------------
    # PRODUCT
    # --------------------------------------------------------

    product_sheet = workbook.create_sheet(
        "Product Analysis"
    )

    products = (
        df.groupby("Product ID")
        .agg(
            Total_Inventory=(
                "Inventory Level",
                "sum"
            ),
            Total_Sales=(
                "Units Sold",
                "sum"
            ),
            Total_Demand=(
                "Demand",
                "sum"
            )
        )
        .reset_index()
    )

    product_sheet.append(
        list(products.columns)
    )

    for row in products.itertuples(
        index=False,
        name=None
    ):

        product_sheet.append(
            list(row)
        )

    # --------------------------------------------------------
    # ALERTS
    # --------------------------------------------------------

    alert_sheet = workbook.create_sheet(
        "Stock Alerts"
    )

    alert_columns = [

        "Product ID",
        "Category",
        "Inventory Level",
        "Demand",
        "Reorder Level",
        "Reorder Quantity",
        "Stock Status",
        "Smart Recommendation"

    ]

    alert_sheet.append(
        alert_columns
    )

    alerts = df[
        df["Smart Recommendation"] !=
        "No Action"
    ][alert_columns]

    for row in alerts.itertuples(
        index=False,
        name=None
    ):

        alert_sheet.append(
            list(row)
        )

    # --------------------------------------------------------
    # MONTHLY
    # --------------------------------------------------------

    monthly_sheet = workbook.create_sheet(
        "Monthly Analysis"
    )

    monthly = (
        df.groupby("Year-Month")
        .agg(
            Inventory_Level=(
                "Inventory Level",
                "mean"
            ),
            Units_Sold=(
                "Units Sold",
                "sum"
            ),
            Demand=(
                "Demand",
                "sum"
            )
        )
        .reset_index()
    )

    monthly_sheet.append(
        list(monthly.columns)
    )

    for row in monthly.itertuples(
        index=False,
        name=None
    ):

        monthly_sheet.append(
            list(row)
        )

    # --------------------------------------------------------
    # FORECAST
    # --------------------------------------------------------

    forecast_sheet = workbook.create_sheet(
        "Forecast"
    )

    forecast = load_forecast()

    if not forecast.empty:

        forecast_sheet.append(
            list(forecast.columns)
        )

        for row in forecast.itertuples(
            index=False,
            name=None
        ):

            forecast_sheet.append(
                list(row)
            )

    # --------------------------------------------------------
    # MODEL COMPARISON
    # --------------------------------------------------------

    model_sheet = workbook.create_sheet(
        "Model Comparison"
    )

    models = load_model_comparison()

    if not models.empty:

        model_sheet.append(
            list(models.columns)
        )

        for row in models.itertuples(
            index=False,
            name=None
        ):

            model_sheet.append(
                list(row)
            )

    # --------------------------------------------------------
    # FORMATTING
    # --------------------------------------------------------

    for ws in workbook.worksheets:

        for cell in ws[1]:

            cell.font = Font(
                bold=True
            )

        for column in ws.columns:

            letter = (
                column[0]
                .column_letter
            )

            max_length = max(
                len(
                    str(cell.value)
                )
                if cell.value is not None
                else 0
                for cell in column
            )

            ws.column_dimensions[
                letter
            ].width = min(
                max_length + 3,
                40
            )

        for row in ws.iter_rows():

            for cell in row:

                cell.alignment = Alignment(
                    vertical="center"
                )

    output = BytesIO()

    workbook.save(output)

    output.seek(0)

    return output


# ============================================================
# REPORT API
# ============================================================

@app.route("/api/generate-report")
def generate_report():

    try:

        df = load_inventory_data()

        report = create_excel_report(
            df
        )

        return send_file(

            report,

            mimetype=
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet",

            as_attachment=True,

            download_name=
            "stockwise_inventory_report.xlsx"
        )

    except Exception as error:

        return jsonify({

            "success": False,

            "error": str(error)

        }), 500


# ============================================================
# CSS
# ============================================================

@app.route("/style.css")
def stylesheet():

    return send_from_directory(
        FRONTEND_DIR,
        "style.css"
    )


# ============================================================
# HEALTH CHECK
# ============================================================

@app.route("/health")
def health():

    return jsonify({

        "status": "online",

        "application":
            "StockWise Smart Inventory Management System"
    })


# ============================================================
# RUN
# ============================================================

if __name__ == "__main__":

    print("=" * 60)
    print(
        "StockWise - Smart Inventory Management System"
    )
    print("=" * 60)

    print(
        "Dashboard: "
        "http://127.0.0.1:5000"
    )

    print(
        "API: "
        "http://127.0.0.1:5000/api/dashboard"
    )

    print("=" * 60)

    app.run(
        host="127.0.0.1",
        port=5000,
        debug=True
    )